#!/usr/bin/env python3
# /// script
# requires-python = ">=3.11"
# dependencies = ["python-docx>=1.1,<2"]
# ///
"""Generate the client-facing Word paper outline from its Markdown authority."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


REPO_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = REPO_ROOT / "docs/paper/paper-outline.md"
DEFAULT_OUTPUT = REPO_ROOT / "output/doc/constructing-screen-time-paper-outline.docx"
ICON_PATH = REPO_ROOT / "web/public/icon-512.png"

NAVY = "1F2A44"
INK = "0F1729"
MUTED = "4B5670"
TEAL = "0E8A82"
ORANGE = "F4B266"
PALE_ORANGE = "FFF4E5"
PALE = "F7F8FB"
BORDER = "DCE1E8"
WHITE = "FFFFFF"


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=80, bottom=70, end=80) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = tc_mar.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def remove_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "nil")


def add_field(paragraph, instruction: str, display_text: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction_text = OxmlElement("w:instrText")
    instruction_text.set(qn("xml:space"), "preserve")
    instruction_text.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = display_text
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction_text, separate, text, end))


def add_hyperlink(paragraph, label: str, target: str) -> None:
    relationship_id = paragraph.part.relate_to(
        target,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), TEAL)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.extend((color, underline))
    text = OxmlElement("w:t")
    text.text = label
    run.extend((run_properties, text))
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_PATTERN = re.compile(r"(\*\*.+?\*\*|`.+?`|\[[^\]]+\]\([^)]+\))")


def add_inline(paragraph, value: str, *, size: float | None = None, color: str | None = None) -> None:
    cursor = 0
    for match in INLINE_PATTERN.finditer(value):
        if match.start() > cursor:
            run = paragraph.add_run(value[cursor : match.start()])
            if size:
                run.font.size = Pt(size)
            if color:
                run.font.color.rgb = rgb(color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Liberation Mono"
            run.font.size = Pt(9 if size is None else min(size, 9))
            run.font.color.rgb = rgb(NAVY)
            r_pr = run._r.get_or_add_rPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "EEF1F5")
            r_pr.append(shd)
        else:
            link = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token)
            assert link
            add_hyperlink(paragraph, link.group(1), link.group(2))
        cursor = match.end()
    if cursor < len(value):
        run = paragraph.add_run(value[cursor:])
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = rgb(color)


def set_style_font(style, name: str, size: float, color: str, bold: bool = False) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = rgb(color)
    style.font.bold = bold
    style._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def configure_styles(document: Document) -> None:
    styles = document.styles

    normal = styles["Normal"]
    set_style_font(normal, "Arial", 10.5, INK)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.12

    title = styles["Title"]
    set_style_font(title, "Arial", 25, NAVY, bold=True)
    title.paragraph_format.space_after = Pt(8)

    subtitle = styles["Subtitle"]
    set_style_font(subtitle, "Arial", 12, TEAL, bold=True)

    for style_name, size, color, before, after in (
        ("Heading 1", 16, NAVY, 14, 6),
        ("Heading 2", 12.5, TEAL, 11, 4),
        ("Heading 3", 10.8, NAVY, 8, 3),
    ):
        style = styles[style_name]
        set_style_font(style, "Arial", size, color, bold=True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    styles["Heading 1"].paragraph_format.page_break_before = True

    caption = styles["Caption"]
    set_style_font(caption, "Arial", 9.5, NAVY, bold=True)
    caption.paragraph_format.space_before = Pt(8)
    caption.paragraph_format.space_after = Pt(4)
    caption.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        set_style_font(style, "Arial", 10.3, INK)
        style.paragraph_format.space_after = Pt(2.5)

    if "Numbered Item" not in styles:
        numbered_item = styles.add_style("Numbered Item", WD_STYLE_TYPE.PARAGRAPH)
    else:
        numbered_item = styles["Numbered Item"]
    set_style_font(numbered_item, "Arial", 10.3, INK)
    numbered_item.paragraph_format.left_indent = Inches(0.24)
    numbered_item.paragraph_format.first_line_indent = Inches(-0.18)
    numbered_item.paragraph_format.space_after = Pt(2.5)

    if "Code Block" not in styles:
        code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = styles["Code Block"]
    set_style_font(code, "Liberation Mono", 8.5, NAVY)
    code.paragraph_format.left_indent = Inches(0.18)
    code.paragraph_format.right_indent = Inches(0.08)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(6)

    if "Callout" not in styles:
        callout = styles.add_style("Callout", WD_STYLE_TYPE.PARAGRAPH)
    else:
        callout = styles["Callout"]
    set_style_font(callout, "Arial", 9.5, INK)
    callout.paragraph_format.left_indent = Inches(0.18)
    callout.paragraph_format.right_indent = Inches(0.08)
    callout.paragraph_format.space_before = Pt(5)
    callout.paragraph_format.space_after = Pt(7)


def set_section_geometry(section, *, landscape: bool = False) -> None:
    section.different_first_page_header_footer = False
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11)
        section.page_height = Inches(8.5)
        section.left_margin = Inches(0.55)
        section.right_margin = Inches(0.55)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)


def configure_header_footer(section) -> None:
    section.header.is_linked_to_previous = True
    section.footer.is_linked_to_previous = True
    if section.header.paragraphs:
        header = section.header.paragraphs[0]
    else:
        header = section.header.add_paragraph()
    if not header.text:
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = header.add_run("CONSTRUCTING SCREEN TIME  |  MILESTONE-ZERO OUTLINE")
        run.font.name = "Arial"
        run.font.size = Pt(7.5)
        run.font.color.rgb = rgb(MUTED)

    if section.footer.paragraphs:
        footer = section.footer.paragraphs[0]
    else:
        footer = section.footer.add_paragraph()
    if not footer.text:
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("2026-08-11  |  Page ")
        run.font.name = "Arial"
        run.font.size = Pt(8)
        run.font.color.rgb = rgb(MUTED)
        add_field(footer, "PAGE", "1")


def add_title_page(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(16)
    if ICON_PATH.exists():
        paragraph.add_run().add_picture(str(ICON_PATH), width=Inches(0.62))

    title = document.add_paragraph(style="Title")
    title.add_run("Constructing Screen Time")
    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.add_run("From Typed Android Events to Cross-Platform Measures")

    rule = document.add_paragraph()
    rule.paragraph_format.space_before = Pt(1)
    rule.paragraph_format.space_after = Pt(14)
    bottom = OxmlElement("w:pBdr")
    border = OxmlElement("w:bottom")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), "18")
    border.set(qn("w:space"), "1")
    border.set(qn("w:color"), TEAL)
    bottom.append(border)
    rule._p.get_or_add_pPr().append(bottom)

    label = document.add_paragraph()
    add_inline(label, "**Detailed manuscript outline**", size=12)

    metadata = document.add_table(rows=4, cols=2)
    metadata.alignment = WD_TABLE_ALIGNMENT.LEFT
    remove_table_borders(metadata)
    values = (
        ("Status", "Milestone-zero authoring plan; headline empirical analyses unrun"),
        ("Paper type", "Measurement methods and empirical validation"),
        ("Prepared", "11 August 2026"),
        ("Primary source", "docs/paper/paper-outline.md"),
    )
    for row, (key, value) in zip(metadata.rows, values, strict=True):
        row.cells[0].width = Inches(1.25)
        row.cells[1].width = Inches(5.7)
        for cell in row.cells:
            set_cell_margins(cell, top=45, bottom=45)
        p0 = row.cells[0].paragraphs[0]
        r0 = p0.add_run(key.upper())
        r0.bold = True
        r0.font.size = Pt(8)
        r0.font.color.rgb = rgb(MUTED)
        p1 = row.cells[1].paragraphs[0]
        add_inline(p1, value, size=9.5, color=INK)

    callout = document.add_paragraph(style="Callout")
    add_inline(
        callout,
        "**Evidence warning.** This is an authoring outline, not a result report. Empty Results "
        "fields are deliberate. The 60-minute mass-conservation equality is a mathematical "
        "identity, not empirical evidence of accurate within-hour reconstruction.",
    )
    p_pr = callout._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), PALE_ORANGE)
    p_pr.append(shading)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), ORANGE)
    borders.append(left)
    p_pr.append(borders)

    document.add_page_break()
    toc_title = document.add_paragraph()
    toc_title.paragraph_format.space_before = Pt(6)
    toc_title.paragraph_format.space_after = Pt(7)
    toc_title.paragraph_format.keep_with_next = True
    toc_title_run = toc_title.add_run("Contents")
    toc_title_run.bold = True
    toc_title_run.font.name = "Arial"
    toc_title_run.font.size = Pt(16)
    toc_title_run.font.color.rgb = rgb(NAVY)
    toc = document.add_paragraph()
    add_field(toc, 'TOC \\o "1-3" \\h \\z \\u', "Right-click and update this field in Word.")
    document.add_page_break()


def add_table(document: Document, rows: list[list[str]], caption: str | None) -> None:
    column_count = max(len(row) for row in rows)
    landscape = column_count >= 6
    if landscape:
        section = document.add_section(WD_SECTION.NEW_PAGE)
        set_section_geometry(section, landscape=True)
        configure_header_footer(section)

    if caption:
        paragraph = document.add_paragraph(style="Caption")
        if caption.startswith("Table 10."):
            paragraph.paragraph_format.page_break_before = True
        add_inline(paragraph, caption)

    table = document.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    font_size = 8.0 if column_count >= 7 else 8.3 if column_count >= 5 else 9.0

    for row_index, row_values in enumerate(rows):
        row = table.rows[row_index]
        prevent_row_split(row)
        if row_index == 0:
            set_repeat_table_header(row)
        for column_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cell)
            value = row_values[column_index] if column_index < len(row_values) else ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(1.5)
            add_inline(paragraph, value, size=font_size)
            for run in paragraph.runs:
                run.font.name = "Arial"
                run.font.size = Pt(font_size)
                if row_index == 0:
                    run.bold = True
                    run.font.color.rgb = rgb(WHITE)
            if row_index == 0:
                set_cell_shading(cell, NAVY)
            elif row_index % 2 == 0:
                set_cell_shading(cell, PALE)

    document.add_paragraph().paragraph_format.space_after = Pt(2)
    if landscape:
        section = document.add_section(WD_SECTION.NEW_PAGE)
        set_section_geometry(section, landscape=False)
        configure_header_footer(section)


def parse_table(lines: list[str]) -> list[list[str]]:
    parsed: list[list[str]] = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if cells and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            continue
        parsed.append(cells)
    return parsed


def add_code_block(document: Document, code_lines: list[str], caption: str | None) -> None:
    if caption:
        paragraph = document.add_paragraph(style="Caption")
        add_inline(paragraph, caption)
    paragraph = document.add_paragraph(style="Code Block")
    paragraph.add_run("\n".join(code_lines))
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "EEF1F5")
    p_pr.append(shading)


def add_markdown_body(document: Document, source: str) -> None:
    lines = source.splitlines()
    start = next(index for index, line in enumerate(lines) if line == "## Abstract")
    index = start
    paragraph_buffer: list[str] = []
    pending_caption: str | None = None

    def flush_paragraph() -> None:
        nonlocal paragraph_buffer
        if paragraph_buffer:
            paragraph = document.add_paragraph()
            add_inline(paragraph, " ".join(value.strip() for value in paragraph_buffer))
            paragraph_buffer = []

    def flush_standalone_caption() -> None:
        nonlocal pending_caption
        if pending_caption:
            paragraph = document.add_paragraph(style="Caption")
            add_inline(paragraph, pending_caption)
            pending_caption = None

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if not stripped:
            flush_paragraph()
            index += 1
            continue

        # Table and code-block captions are emitted with the object they label. A figure-plan
        # caption is commonly followed by prose or bullets, so it must be emitted before that
        # content instead of being overwritten by the next caption.
        if pending_caption and not stripped.startswith(("|", "```")):
            flush_standalone_caption()

        heading = re.match(r"^(#{2,4})\s+(.+)$", stripped)
        if heading:
            flush_paragraph()
            level = len(heading.group(1)) - 1
            text = heading.group(2)
            if text.startswith("Table ") or text.startswith("Figure "):
                pending_caption = text
            else:
                paragraph = document.add_paragraph(style=f"Heading {level}")
                if text == "Abstract":
                    paragraph.paragraph_format.page_break_before = False
                add_inline(paragraph, text)
            index += 1
            continue

        if stripped.startswith("|"):
            flush_paragraph()
            table_lines: list[str] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_table(document, parse_table(table_lines), pending_caption)
            pending_caption = None
            continue

        if stripped.startswith("```"):
            flush_paragraph()
            index += 1
            code_lines: list[str] = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index])
                index += 1
            index += 1
            add_code_block(document, code_lines, pending_caption)
            pending_caption = None
            continue

        bullet = re.match(r"^\s*-\s+(.+)$", line)
        numbered = re.match(r"^\s*(\d+)\.\s+(.+)$", line)
        if bullet or numbered:
            flush_paragraph()
            match = bullet or numbered
            assert match
            content = match.group(1) if bullet else match.group(2)
            item_number = None if bullet else match.group(1)
            index += 1
            continuations: list[str] = []
            while index < len(lines):
                continuation = lines[index]
                if not continuation.strip() or re.match(r"^\s*(?:-|\d+\.)\s+", continuation):
                    break
                if continuation.startswith(" ") and not continuation.lstrip().startswith(("#", "|", "```")):
                    continuations.append(continuation.strip())
                    index += 1
                else:
                    break
            full_text = " ".join((content, *continuations))
            style = "List Bullet" if bullet else "Numbered Item"
            paragraph = document.add_paragraph(style=style)
            if item_number is not None:
                full_text = f"{item_number}. {full_text}"
            add_inline(paragraph, full_text)
            continue

        if stripped.startswith(">"):
            flush_paragraph()
            quote_lines: list[str] = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                quote_lines.append(lines[index].strip().lstrip(">").strip())
                index += 1
            paragraph = document.add_paragraph(style="Callout")
            add_inline(paragraph, " ".join(quote_lines))
            p_pr = paragraph._p.get_or_add_pPr()
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), PALE_ORANGE)
            p_pr.append(shading)
            continue

        paragraph_buffer.append(stripped)
        index += 1

    flush_paragraph()


def set_document_properties(document: Document) -> None:
    properties = document.core_properties
    properties.title = "Constructing Screen Time: From Typed Android Events to Cross-Platform Measures"
    properties.subject = "Detailed manuscript outline"
    properties.author = "Chronicle research team"
    properties.keywords = "screen time, Android, iOS, measurement, temporal resolution, validation"
    properties.comments = "Generated from docs/paper/paper-outline.md"


def build_document(source_path: Path, output_path: Path) -> None:
    source = source_path.read_text(encoding="utf-8")
    if "## 1. Introduction" not in source or "### 2.11 Study 3B" not in source:
        raise SystemExit("source does not contain the expected manuscript and resolution-study sections")

    document = Document()
    set_document_properties(document)
    configure_styles(document)
    set_section_geometry(document.sections[0], landscape=False)
    document.sections[0].different_first_page_header_footer = True
    configure_header_footer(document.sections[0])
    add_title_page(document)
    add_markdown_body(document, source)

    settings = document.settings._element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    arguments = parser.parse_args()
    build_document(arguments.source.resolve(), arguments.output.resolve())
    print(arguments.output.resolve())


if __name__ == "__main__":
    main()
